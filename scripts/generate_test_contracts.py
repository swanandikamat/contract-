import os
from pathlib import Path
import docx

BASE_DIR = Path(__file__).resolve().parent.parent
TEST_CLEAN_DIR = BASE_DIR / "data" / "test_contracts" / "clean"
TEST_PROB_DIR = BASE_DIR / "data" / "test_contracts" / "prob"

TEST_CLEAN_DIR.mkdir(parents=True, exist_ok=True)
TEST_PROB_DIR.mkdir(parents=True, exist_ok=True)

def create_docx(filename: Path, title: str, ref: str, sections: list):
    doc = docx.Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Contract Reference: {ref}")
    doc.add_paragraph(
        "This Agreement (\"Agreement\") is entered into by and between Syngenta Crop Protection AG (\"Syngenta\") "
        "and Enterprise Solutions Inc. (\"Provider\"), effective as of October 1, 2026."
    )
    
    doc.add_heading("Recitals", level=1)
    doc.add_paragraph("WHEREAS, Syngenta requires modern digital enterprise cloud and consulting services; and")
    doc.add_paragraph("WHEREAS, Provider possesses the requisite technical expertise and capacity to perform such services under balanced commercial terms.")
    
    for section_title, section_text in sections:
        doc.add_heading(section_title, level=2)
        doc.add_paragraph(section_text)
        
    doc.save(str(filename))
    print(f"Created docx: {filename}")

# 1. TEST-CLEAN-01: Cloud Services Agreement
clean_01_sections = [
    ("1. Scope of Cloud Services", "Provider shall make available enterprise cloud platform access to Syngenta in accordance with the Service Level Agreement (SLA) specified in Schedule A."),
    ("2. Payment Terms & Invoicing", "Syngenta shall pay undisputed invoices within thirty (30) days of receipt. Invoices shall clearly itemize subscription tiers and usage metrics."),
    ("3. Limitation of Liability", "Except for breaches of confidentiality or willful misconduct, each party's maximum aggregate liability under this Agreement shall be limited to the total fees paid or payable by Syngenta in the twelve (12) months preceding the claim."),
    ("4. Mutual Indemnification", "Each party agrees to defend, indemnify, and hold harmless the other party from and against third-party claims arising from gross negligence, willful misconduct, or infringement of third-party intellectual property rights."),
    ("5. Term & Notice of Renewal", "This Agreement shall remain in effect for one (1) year. Either party may terminate or opt out of automatic renewal by providing written notice at least thirty (30) days prior to the expiration of the current term."),
    ("6. Intellectual Property & Data Ownership", "Syngenta retains all right, title, and interest in and to its data, trademarks, and proprietary software. Provider acquires no rights to Syngenta data except as expressly needed to render the services."),
    ("7. Governing Law & Dispute Resolution", "This Agreement shall be governed by and construed in accordance with the laws of Switzerland. Any disputes shall be resolved through good-faith negotiations, and if necessary, binding arbitration under ICC rules in Zurich.")
]

# 2. TEST-CLEAN-02: Professional Consulting Agreement
clean_02_sections = [
    ("1. Consulting Services & Milestone Deliverables", "Consultant will perform digital transformation advisory services as detailed in Statement of Work #1. Deliverables will be reviewed by Syngenta within fifteen (15) business days."),
    ("2. Fees and Reimbursable Expenses", "Services will be billed on a milestone completion basis. Reasonable, pre-approved travel and accommodation expenses will be reimbursed at cost without markup upon submission of receipts."),
    ("3. Standard Risk & Liability Cap", "To the maximum extent permitted by applicable law, aggregate liability of either party for claims connected with this Agreement shall not exceed two times (2x) the total contract value."),
    ("4. Confidentiality & Non-Disclosure", "Both parties shall maintain strict confidentiality over proprietary technical and business information disclosed during the term of this Agreement for a period of three (3) years post-termination."),
    ("5. Termination for Convenience", "Either party may terminate this Agreement without cause by giving thirty (30) days' prior written notice to the other party. Syngenta shall pay for all satisfactory services rendered up to the effective date of termination."),
    ("6. Compliance & Data Protection", "Consultant agrees to comply with all applicable data protection laws, including GDPR and local Swiss data regulations, when processing personal data on behalf of Syngenta.")
]

# 3. TEST-PROB-01: Problematic Cloud Services Agreement
prob_01_sections = [
    ("1. Cloud Platform Access & SLA", "Provider will provide software access on an 'as-is' basis without guaranteed uptime or service levels. Provider reserves the right to modify features at any time without prior notification."),
    ("2. Pricing, Taxes & Currency Adjustment Trap", "Syngenta shall settle all invoices within seven (7) days. Provider reserves the unilateral right to apply an automatic 15% quarterly rate hike and levy a 5% monthly penalty for any payment delay beyond 7 days."),
    ("3. Uncapped Consequential Liability & One-Sided Cap", "Provider's aggregate liability under this Agreement is strictly capped at $100 total. Syngenta shall be liable for all direct, indirect, special, punitive, and consequential damages of any kind incurred by Provider without any limit whatsoever."),
    ("4. One-Sided Indemnification", "Syngenta shall fully defend, indemnify, and hold harmless Provider against all claims, losses, legal costs, and liabilities arising out of or related to Provider's provision of services or any operational disruption."),
    ("5. Auto-Renewal Trap & Excessive Lock-In", "This Agreement automatically renews for successive three (3) year terms unless Syngenta delivers written notice of non-renewal exactly between 180 and 175 days prior to expiration. Failure to meet this precise 5-day notice window results in irrevocable commitment to the full 3-year term."),
    ("6. Data Expropriation & IP Transfer", "Provider retains sole ownership of all customer data, derivative analytics, machine learning models, and system enhancements generated through Syngenta's use of the platform. Syngenta forfeits all IP claims upon contract execution.")
]

# 4. TEST-PROB-02: Problematic Professional Consulting Agreement
prob_02_sections = [
    ("1. Vague Scope & Upfront Advance Forfeiture", "Consultant shall provide strategic advisory services at its sole discretion. Syngenta shall pay 100% of the total estimated fee upfront as a non-refundable advance regardless of deliverable completion or quality."),
    ("2. Uncapped Breach Liability & Strict Penalty", "Syngenta shall indemnify Consultant against any lawsuit or claim brought by any third party. Syngenta assumes full, uncapped liability for any indirect or consequential damages arising from any project delay."),
    ("3. Overbroad Non-Compete & Restrictive Covenant", "Syngenta agrees that for a period of ten (10) years following termination, it shall not engage, hire, or partner with any technology consulting firm operating anywhere globally in the agricultural or chemical domain."),
    ("4. One-Sided Immediate Termination & Restitution", "Consultant may terminate this Agreement at any time with zero days' notice without refunding any fees. Syngenta may not terminate this Agreement for any reason prior to full contract expiration."),
    ("5. Foreign Jurisdiction & One-Sided Fee Shifting", "This Agreement shall be exclusively governed by the laws of the Republic of Vanuatu. Syngenta consents to the exclusive jurisdiction of the court of Vanuatu and agrees to pay all attorney fees incurred by Consultant in any dispute.")
]

def main():
    create_docx(TEST_CLEAN_DIR / "TEST-CLEAN-01_Cloud_Services_Agreement.docx", "Cloud Services Agreement", "TEST-CLEAN-01", clean_01_sections)
    create_docx(TEST_CLEAN_DIR / "TEST-CLEAN-02_Professional_Consulting_Agreement.docx", "Professional Consulting Agreement", "TEST-CLEAN-02", clean_02_sections)
    create_docx(TEST_PROB_DIR / "TEST-PROB-01_Cloud_Services_Agreement.docx", "Problematic Cloud Services Agreement", "TEST-PROB-01", prob_01_sections)
    create_docx(TEST_PROB_DIR / "TEST-PROB-02_Professional_Consulting_Agreement.docx", "Problematic Consulting Agreement", "TEST-PROB-02", prob_02_sections)
    print("All synthetic test documents generated successfully!")

if __name__ == "__main__":
    main()
